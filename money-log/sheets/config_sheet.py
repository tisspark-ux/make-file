"""설정 시트 생성 모듈"""

from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from version import VERSION


# ── 스타일 헬퍼 ──────────────────────────────────────────────
def _side():
    return Side(style="thin", color="CCCCCC")

def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, size=11, color="000000"):
    return Font(bold=bold, size=size, color=color)

def _align(h="center", v="center"):
    return Alignment(horizontal=h, vertical=v, wrap_text=True)


# ── 카테고리 매핑 테이블 ──────────────────────────────────────
# (뱅샐 대분류, 뱅샐 소분류) → 가계부 대분류
# 소분류가 "*" 이면 해당 대분류 전체에 적용
CATEGORY_MAP = [
    # 뱅샐 대분류     뱅샐 소분류         가계부 대분류
    ("고정비",        "교통비",           "고정비"),
    ("고정비",        "휴대폰",           "고정비"),
    ("고정비",        "인터넷",           "고정비"),
    ("고정비",        "보험",             "고정비"),
    ("고정비",        "구독",             "고정비"),
    ("고정비",        "*",               "고정비"),
    ("생활비",        "외식",             "생활비"),
    ("생활비",        "식료품",           "생활비"),
    ("생활비",        "생활용품",         "생활비"),
    ("생활비",        "*",               "생활비"),
    ("의료/건강",     "*",               "의료/건강"),
    ("모임",          "*",               "모임"),
    ("세금",          "*",               "세금"),
    ("대출상환",      "*",               "대출상환"),
    ("부동산",        "*",               "부동산"),
    ("Tiss",         "*",               "Tiss"),
    ("JM",           "*",               "JM"),
    ("금융수입",      "*",               "금융수입"),
    ("기타",          "*",               "기타"),
]


# ── 시트 생성 ─────────────────────────────────────────────────
def build(wb, year: int = 2026):
    ws = wb.create_sheet("설정", 0)
    ws.sheet_view.showGridLines = False

    # 열 너비
    col_widths = {1: 4, 2: 18, 3: 18, 4: 14, 5: 14, 6: 4}
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for r in range(1, 60):
        ws.row_dimensions[r].height = 20

    # ── 기본 설정 ────────────────────────────────────────────
    _section_title(ws, 2, 2, "기본 설정")

    settings = [
        ("연도",    year,            "이 셀만 바꾸면 전체 시트 자동 업데이트"),
        ("담당자1", "Tiss",          ""),
        ("담당자2", "JM",            ""),
        ("버전",    f"v{VERSION}",   "마이너 +0.01 / 메이저 +1.00"),
    ]
    for i, (label, value, note) in enumerate(settings, start=3):
        _label_cell(ws, i, 2, label)
        c = ws.cell(row=i, column=3, value=value)
        c.font = _font(bold=(i == 3), size=12 if i == 3 else 11,
                       color="C00000" if i == 3 else "000000")
        c.alignment = _align()
        c.border = _border()
        if note:
            ws.cell(row=i, column=4, value=note).font = Font(size=9, color="888888",
                                                              italic=True)

    # ── 카테고리 매핑 테이블 ──────────────────────────────────
    _section_title(ws, 8, 2, "뱅샐 카테고리 매핑")

    headers = ["뱅샐 대분류", "뱅샐 소분류", "가계부 대분류"]
    for ci, h in enumerate(headers, start=2):
        c = ws.cell(row=9, column=ci, value=h)
        c.fill = _fill("1F4E79")
        c.font = _font(bold=True, color="FFFFFF")
        c.alignment = _align()
        c.border = _border()

    for ri, (banksalad_main, banksalad_sub, budget_main) in enumerate(CATEGORY_MAP, start=10):
        sub_display = banksalad_sub if banksalad_sub != "*" else "(전체)"
        row_data = [banksalad_main, sub_display, budget_main]
        fill_color = "EBF3FB" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row_data, start=2):
            c = ws.cell(row=ri, column=ci, value=val)
            c.fill = _fill(fill_color)
            c.alignment = _align()
            c.border = _border()

    # 안내 메모
    note_row = 10 + len(CATEGORY_MAP) + 1
    note = ws.cell(row=note_row, column=2,
                   value="※ 매핑되지 않은 항목은 자동으로 '기타'로 분류됩니다.")
    note.font = Font(size=9, color="888888", italic=True)
    ws.merge_cells(f"B{note_row}:E{note_row}")

    return ws


# ── 내부 헬퍼 ─────────────────────────────────────────────────
def _section_title(ws, row, col, title):
    c = ws.cell(row=row, column=col, value=title)
    c.font = _font(bold=True, size=12, color="1F4E79")
    c.fill = _fill("D6E4F0")
    c.alignment = _align(h="left")
    ws.merge_cells(f"{get_column_letter(col)}{row}:{get_column_letter(col+3)}{row}")
    c.border = _border()

def _label_cell(ws, row, col, label):
    c = ws.cell(row=row, column=col, value=label)
    c.font = _font(bold=True, color="FFFFFF")
    c.fill = _fill("2E75B6")
    c.alignment = _align()
    c.border = _border()
