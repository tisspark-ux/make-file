"""대시보드 시트 생성 모듈 — 이번달 현황 한눈에 보기"""

from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.utils import get_column_letter
from .budget_sheet import BUDGET_ITEMS


def _side():
    return Side(style="thin", color="CCCCCC")
def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)
def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)
def _font(bold=False, size=11, color="000000"):
    return Font(bold=bold, size=size, color=color)
def _align(h="center", v="center"):
    return Alignment(horizontal=h, vertical=v, wrap_text=False)


# 월별 시트에서 예산/실제 컬럼 위치 (monthly_sheet.py 기준)
_MONTHLY_COL_BASE   = "C"  # 예산기본
_MONTHLY_COL_TOTAL  = "E"  # 예산합계
_MONTHLY_COL_ACTUAL = "F"  # 실제지출
_MONTHLY_DATA_START = 5    # 월별 시트 데이터 시작 행


def _indirect(col, row):
    """현재 월 시트를 동적으로 참조하는 INDIRECT 수식"""
    return (f'=IFERROR(INDIRECT("\'"\u0026'
            f'TEXT(MONTH(TODAY()),"00")\u0026'
            f'"월!{col}{row}"),0)')


def build(wb):
    ws = wb.create_sheet("대시보드", 3)   # 뱅샐입력 바로 다음
    ws.sheet_view.showGridLines = False
    ws.sheet_view.tabSelected = True     # 파일 열 때 이 시트로

    col_widths = {1: 4, 2: 18, 3: 18, 4: 18, 5: 14, 6: 14, 7: 14, 8: 4}
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    row = 1

    # ── 제목 ─────────────────────────────────────────────────
    ws.row_dimensions[row].height = 32
    title = ws.cell(row=row, column=2,
                    value='=TEXT(TODAY(),"YYYY년 MM월")&" 가계부 현황"')
    title.font = _font(bold=True, size=16, color="1F4E79")
    title.alignment = _align(h="left")
    ws.merge_cells("B1:G1")

    # 갱신 기준
    row = 2
    ws.row_dimensions[row].height = 14
    sub = ws.cell(row=row, column=2,
                  value='="기준: "& TODAY() &"  ·  "& DAYS(EOMONTH(TODAY(),0),TODAY()) &"일 남음"')
    sub.font = Font(size=9, color="888888", italic=True)
    sub.alignment = _align(h="left")
    ws.merge_cells("B2:G2")

    row = 4

    # ── 요약 바 ───────────────────────────────────────────────
    ws.row_dimensions[row].height = 14
    ws.cell(row=row, column=2, value="이번달 수입").font = _font(size=9, color="888888")
    ws.cell(row=row, column=3, value="이번달 지출").font = _font(size=9, color="888888")
    ws.cell(row=row, column=4, value="잔액").font = _font(size=9, color="888888")
    ws.cell(row=row, column=5, value="예산합계").font = _font(size=9, color="888888")
    ws.cell(row=row, column=6, value="달성율").font = _font(size=9, color="888888")
    ws.cell(row=row, column=7, value="하루 평균 지출").font = _font(size=9, color="888888")
    for ci in range(2, 8):
        ws.cell(row=row, column=ci).alignment = _align()

    row = 5
    ws.row_dimensions[row].height = 30
    summary_vals = [
        ('=IFERROR(INDIRECT("\'"\u0026TEXT(MONTH(TODAY()),"00")\u0026"월!B2"),0)', "375623", '#,##0"원"'),
        ('=IFERROR(INDIRECT("\'"\u0026TEXT(MONTH(TODAY()),"00")\u0026"월!D2"),0)', "C00000", '#,##0"원"'),
        ('=IFERROR(INDIRECT("\'"\u0026TEXT(MONTH(TODAY()),"00")\u0026"월!F2"),0)', "1F4E79", '#,##0"원"'),
        (f'=IFERROR(INDIRECT("\'"\u0026TEXT(MONTH(TODAY()),"00")\u0026"월!E{_MONTHLY_DATA_START + len(BUDGET_ITEMS)}"),0)',
         "7030A0", '#,##0"원"'),
        (f'=IFERROR(INDIRECT("\'"\u0026TEXT(MONTH(TODAY()),"00")\u0026"월!D2")/INDIRECT("\'"\u0026TEXT(MONTH(TODAY()),"00")\u0026"월!E{_MONTHLY_DATA_START + len(BUDGET_ITEMS)}"),0)',
         "4472C4", '0%'),
        ('=IFERROR(INDIRECT("\'"\u0026TEXT(MONTH(TODAY()),"00")\u0026"월!D2")/(DAY(TODAY())),0)',
         "404040", '#,##0"원"'),
    ]
    for ci, (val, color, fmt) in enumerate(summary_vals, start=2):
        c = ws.cell(row=row, column=ci, value=val)
        c.fill = _fill(color)
        c.font = _font(bold=True, size=13, color="FFFFFF")
        c.alignment = _align()
        c.number_format = fmt
        c.border = _border()

    row = 7

    # ── 예산 vs 실제 테이블 ──────────────────────────────────
    ws.row_dimensions[row].height = 22
    headers = ["카테고리", "세부항목", "예산", "실제지출", "차액", "달성율"]
    hcolors = ["1F4E79", "1F4E79", "375623", "C00000", "7030A0", "7030A0"]
    for ci, (h, hc) in enumerate(zip(headers, hcolors), start=2):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = _fill(hc)
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.alignment = _align()
        c.border = _border()

    row = 8
    data_start = row
    for i, (main_cat, sub_cat, *_) in enumerate(BUDGET_ITEMS):
        r = row + i
        monthly_row = _MONTHLY_DATA_START + i
        fill_c = "F2F2F2" if i % 2 == 0 else "FFFFFF"
        ws.row_dimensions[r].height = 20

        # 카테고리 / 세부항목
        for ci, val in [(2, main_cat), (3, sub_cat)]:
            c = ws.cell(row=r, column=ci, value=val)
            c.fill = _fill(fill_c)
            c.alignment = _align()
            c.border = _border()

        # 예산 (E열)
        c_budget = ws.cell(row=r, column=4,
                           value=_indirect(_MONTHLY_COL_TOTAL, monthly_row))
        c_budget.number_format = '#,##0"원"'

        # 실제지출 (F열)
        c_actual = ws.cell(row=r, column=5,
                           value=_indirect(_MONTHLY_COL_ACTUAL, monthly_row))
        c_actual.number_format = '#,##0"원"'

        # 차액
        ws.cell(row=r, column=6,
                value=f"=D{r}-E{r}").number_format = '#,##0"원"'

        # 달성율
        ws.cell(row=r, column=7,
                value=f"=IFERROR(E{r}/D{r},0)").number_format = '0%'

        for ci in range(4, 8):
            ws.cell(row=r, column=ci).fill = _fill(fill_c)
            ws.cell(row=r, column=ci).alignment = _align()
            ws.cell(row=r, column=ci).border = _border()

    # ── 합계 행 ──────────────────────────────────────────────
    total_r = data_start + len(BUDGET_ITEMS)
    ws.row_dimensions[total_r].height = 22
    for ci in range(2, 8):
        c = ws.cell(row=total_r, column=ci)
        c.fill = _fill("1F4E79")
        c.font = _font(bold=True, color="FFFFFF")
        c.alignment = _align()
        c.border = _border()
    ws.cell(row=total_r, column=2, value="합계")
    ws.merge_cells(f"B{total_r}:C{total_r}")
    for ci, col_l in [(4, "D"), (5, "E"), (6, "F")]:
        ws.cell(row=total_r, column=ci,
                value=f"=SUM({col_l}{data_start}:{col_l}{total_r-1})").number_format = '#,##0"원"'
    ws.cell(row=total_r, column=7,
            value=f"=IFERROR(E{total_r}/D{total_r},0)").number_format = '0%'

    # ── ① 과소비 조건부 서식 (달성율 > 100%) ─────────────────
    over_fill = PatternFill(start_color="FFD0D0", end_color="FFD0D0", fill_type="solid")
    ws.conditional_formatting.add(
        f"B{data_start}:G{total_r - 1}",
        FormulaRule(formula=[f"=$E{data_start}>$D{data_start}"], fill=over_fill)
    )

    # ── 달성율 컬럼 데이터 바 ─────────────────────────────────
    ws.conditional_formatting.add(
        f"G{data_start}:G{total_r - 1}",
        DataBarRule(start_type="num", start_value=0,
                    end_type="num", end_value=1,
                    color="4472C4")
    )

    ws.freeze_panes = "B8"
    return ws
