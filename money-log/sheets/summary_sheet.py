"""요약 시트 생성 모듈"""

from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.marker import Marker
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from .budget_sheet import BUDGET_ITEMS
from .banksalad_sheet import DATA_LAST_ROW

# 편차 기준: 평균 대비 이 비율(%) 초과 시 비고 입력 촉구
DEVIATION_THRESHOLD = 0.20  # 20%


def _side():
    return Side(style="thin", color="E2E8F0")

def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, size=11, color="000000"):
    return Font(name="Calibri", bold=bold, size=size, color=color)

def _align(h="center", v="center"):
    return Alignment(horizontal=h, vertical=v)


def build(wb, year: int = 2026):
    ws = wb.create_sheet("요약")
    ws.sheet_view.showGridLines = False

    # B~J: 표시 컬럼 / K·L: 차트 평균선 헬퍼 (작은 폭)
    col_widths = {
        1: 4,   # A: 여백
        2: 10,  # B: 월
        3: 16,  # C: 수입
        4: 16,  # D: 지출
        5: 16,  # E: 잔액
        6: 16,  # F: 예산
        7: 10,  # G: 예산대비
        8: 10,  # H: 수입 편차%
        9: 10,  # I: 지출 편차%
        10: 32, # J: 비고 (사유 입력)
        11: 2,  # K: 수입 평균 (차트 전용, 숨김 수준)
        12: 2,  # L: 지출 평균 (차트 전용, 숨김 수준)
    }
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for r in range(1, 60):
        ws.row_dimensions[r].height = 20

    # ── 제목 ─────────────────────────────────────────────────
    title = ws.cell(row=1, column=2, value=f'=설정!C3&"년 가계부 요약"')
    title.font = Font(name="Calibri Light", bold=False, size=18, color="1A2B4A")
    title.alignment = _align(h="left")
    ws.merge_cells("B1:J1")
    ws.row_dimensions[1].height = 30

    # ── 월별 테이블 ───────────────────────────────────────────
    section_row = 3
    _section_title(ws, section_row, "월별 수입 / 지출 / 잔액")

    headers = ["월", "수입", "지출", "잔액", "예산", "예산대비",
               "수입편차", "지출편차", "비고 (이달 특이사항)"]
    header_colors = ["1A2B4A", "1E8449", "C0392B", "1A2B4A",
                     "2471A3", "2471A3", "2471A3", "2471A3", "2471A3"]
    header_row = section_row + 1
    ws.row_dimensions[header_row].height = 22
    for ci, (h, hc) in enumerate(zip(headers, header_colors), start=2):
        c = ws.cell(row=header_row, column=ci, value=h)
        c.fill = _fill(hc)
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.alignment = _align()
        c.border = _border()

    # ── 비고 안내 텍스트 (헤더 아래 작은 글씨) ───────────────
    note_hint_row = header_row + 1
    hint = ws.cell(row=note_hint_row, column=10,
                   value="← 수입·지출이 평균 대비 20% 초과 시 강조. 사유를 입력하세요.")
    hint.font = Font(size=8, color="888888", italic=True)
    hint.alignment = _align(h="left")
    ws.row_dimensions[note_hint_row].height = 14

    # ── 데이터 행 (1~12월) ───────────────────────────────────
    data_start = note_hint_row + 1  # row 6
    budget_total_row_in_monthly = 4 + len(BUDGET_ITEMS) + 1  # 월별 시트의 합계 행

    for month in range(1, 13):
        r = data_start + month - 1
        ws.row_dimensions[r].height = 21
        sheet_name = f"{month:02d}월"
        fill_color = "F4F8FC" if month % 2 == 0 else "FFFFFF"

        # B: 월
        c = ws.cell(row=r, column=2, value=f"{month}월")
        c.fill = _fill(fill_color)
        c.font = _font(bold=True)
        c.alignment = _align()
        c.border = _border()

        # C: 수입
        ws.cell(row=r, column=3,
                value=f"='{sheet_name}'!B2").number_format = '#,##0'
        # D: 지출
        ws.cell(row=r, column=4,
                value=f"='{sheet_name}'!D2").number_format = '#,##0'
        # E: 잔액
        ws.cell(row=r, column=5,
                value=f"='{sheet_name}'!F2").number_format = '#,##0'
        # F: 예산합계
        ws.cell(row=r, column=6,
                value=f"=IFERROR('{sheet_name}'!E{budget_total_row_in_monthly},0)").number_format = '#,##0'
        # G: 예산 달성율
        ws.cell(row=r, column=7,
                value=f"=IFERROR(D{r}/F{r},0)").number_format = '0%'

        # H: 수입 편차% = (이번달 수입 - 평균) / 평균
        avg_income_range = f"C${data_start}:C${data_start+11}"
        ws.cell(row=r, column=8,
                value=f"=IFERROR((C{r}-AVERAGEIF({avg_income_range},\">0\"))"
                      f"/ABS(AVERAGEIF({avg_income_range},\">0\")),0)"
                ).number_format = '+0%;-0%;0%'

        # I: 지출 편차% = (이번달 지출 - 평균) / 평균
        avg_expense_range = f"D${data_start}:D${data_start+11}"
        ws.cell(row=r, column=9,
                value=f"=IFERROR((D{r}-AVERAGEIF({avg_expense_range},\">0\"))"
                      f"/ABS(AVERAGEIF({avg_expense_range},\">0\")),0)"
                ).number_format = '+0%;-0%;0%'

        # J: 비고 (사용자 입력 — 빈 셀로 놔둠)
        ws.cell(row=r, column=10, value="")

        # K: 수입 평균 (차트 평균선 헬퍼)
        ws.cell(row=r, column=11,
                value=f"=IFERROR(AVERAGEIF({avg_income_range},\">0\"),0)"
                ).number_format = '#,##0'
        # L: 지출 평균 (차트 평균선 헬퍼)
        ws.cell(row=r, column=12,
                value=f"=IFERROR(AVERAGEIF({avg_expense_range},\">0\"),0)"
                ).number_format = '#,##0'

        for ci in range(3, 13):
            c = ws.cell(row=r, column=ci)
            c.fill = _fill(fill_color)
            c.alignment = _align()
            c.border = _border()

    # ── 연간 합계 행 ─────────────────────────────────────────
    total_r = data_start + 12
    ws.row_dimensions[total_r].height = 22
    ws.cell(row=total_r, column=2, value="연간 합계")
    for ci, col_l in [(3, "C"), (4, "D"), (5, "E")]:
        ws.cell(row=total_r, column=ci,
                value=f"=SUM({col_l}{data_start}:{col_l}{total_r-1})").number_format = '#,##0'
    for ci in range(2, 11):
        c = ws.cell(row=total_r, column=ci)
        c.fill = _fill("1A2B4A")
        c.font = _font(bold=True, color="FFFFFF")
        c.alignment = _align()
        c.border = _border()

    # ── 조건부 서식 ───────────────────────────────────────────
    data_range_rows = f"{data_start}:{data_start + 11}"

    # ① 수입 이상: |수입편차| > 20% → 수입 셀 노랑
    income_over_fill = PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid")
    ws.conditional_formatting.add(
        f"C{data_start}:C{data_start+11}",
        FormulaRule(formula=[f"=ABS(H{data_start})>{DEVIATION_THRESHOLD}"],
                    fill=income_over_fill)
    )
    # ② 지출 이상: |지출편차| > 20% → 지출 셀 연한 주황
    expense_over_fill = PatternFill(start_color="FFE4C4", end_color="FFE4C4", fill_type="solid")
    ws.conditional_formatting.add(
        f"D{data_start}:D{data_start+11}",
        FormulaRule(formula=[f"=ABS(I{data_start})>{DEVIATION_THRESHOLD}"],
                    fill=expense_over_fill)
    )
    # ③ 비고 미입력 + 편차 초과 → 비고 셀 연한 분홍 (입력 촉구)
    note_prompt_fill = PatternFill(start_color="FFD0D0", end_color="FFD0D0", fill_type="solid")
    ws.conditional_formatting.add(
        f"J{data_start}:J{data_start+11}",
        FormulaRule(
            formula=[f"=AND(J{data_start}=\"\","
                     f"OR(ABS(H{data_start})>{DEVIATION_THRESHOLD},"
                     f"ABS(I{data_start})>{DEVIATION_THRESHOLD}))"],
            fill=note_prompt_fill
        )
    )
    # ④ 비고 입력됨 → 비고 셀 연한 초록 (확인)
    note_done_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
    ws.conditional_formatting.add(
        f"J{data_start}:J{data_start+11}",
        FormulaRule(
            formula=[f"=J{data_start}<>\"\""],
            fill=note_done_fill
        )
    )

    # ── 차트: 수입/지출 바 + 평균선 오버레이 ────────────────
    chart_row = total_r + 2
    _add_chart_with_avg(ws, data_start, chart_row)

    # ── 특이사항 목록 (비고 입력된 달만 자동 표시) ───────────
    events_row = chart_row + 23
    _add_events_section(ws, data_start, events_row)

    # ── ⑤ 미분류 거래 감지 ───────────────────────────────────
    unclassified_row = events_row + 18
    _add_unclassified_section(ws, unclassified_row)

    return ws


def _add_chart_with_avg(ws, data_start, anchor_row):
    """수입/지출 막대 + 평균선 혼합 차트"""
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "월별 수입 / 지출  (점선: 연평균)"
    bar.y_axis.title = "금액 (원)"
    bar.x_axis.title = "월"
    bar.style = 2
    bar.width = 22
    bar.height = 13
    bar.y_axis.majorGridlines = None   # 수평 그리드라인 제거

    # 수입·지출 바 데이터
    income_ref  = Reference(ws, min_col=3, min_row=data_start, max_row=data_start + 11)
    expense_ref = Reference(ws, min_col=4, min_row=data_start, max_row=data_start + 11)
    bar.add_data(income_ref,  titles_from_data=False)
    bar.add_data(expense_ref, titles_from_data=False)
    bar.series[0].title = SeriesLabel(v="수입")
    bar.series[1].title = SeriesLabel(v="지출")
    bar.series[0].graphicalProperties.solidFill = "1E8449"
    bar.series[1].graphicalProperties.solidFill = "C0392B"

    # 평균선 라인 차트 (K=수입평균, L=지출평균)
    line = LineChart()
    avg_income_ref  = Reference(ws, min_col=11, min_row=data_start, max_row=data_start + 11)
    avg_expense_ref = Reference(ws, min_col=12, min_row=data_start, max_row=data_start + 11)
    line.add_data(avg_income_ref,  titles_from_data=False)
    line.add_data(avg_expense_ref, titles_from_data=False)
    line.series[0].title = SeriesLabel(v="수입 평균")
    line.series[1].title = SeriesLabel(v="지출 평균")

    # 점선 스타일
    for i, (color, dash) in enumerate([("1E8449", "dash"), ("C0392B", "dash")]):
        line.series[i].graphicalProperties.line.solidFill = color
        line.series[i].graphicalProperties.line.dashDot = dash
        line.series[i].graphicalProperties.line.width = 20000  # 1.5pt
        m = Marker()
        m.symbol = "none"
        line.series[i].marker = m
        line.series[i].smooth = True

    # 카테고리 (월 라벨)
    cats = Reference(ws, min_col=2, min_row=data_start, max_row=data_start + 11)
    bar.set_categories(cats)
    line.set_categories(cats)

    # 바 + 라인 합체
    bar += line
    ws.add_chart(bar, f"B{anchor_row}")


def _add_events_section(ws, data_start, anchor_row):
    """비고가 입력된 달의 특이사항 목록 (FILTER 자동 표시)"""
    _section_title(ws, anchor_row, "특이사항 목록  (비고 입력 시 자동 표시)")

    headers = ["월", "수입편차", "지출편차", "비고"]
    header_colors = ["1A2B4A", "2471A3", "2471A3", "2471A3"]
    hrow = anchor_row + 1
    ws.row_dimensions[hrow].height = 22
    for ci, (h, hc) in enumerate(zip(headers, header_colors), start=2):
        c = ws.cell(row=hrow, column=ci, value=h)
        c.fill = _fill(hc)
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.alignment = _align()
        c.border = _border()

    # FILTER: 비고(J열)가 비어있지 않은 행만 추출
    data_row = anchor_row + 2
    ws.cell(row=data_row, column=2,
            value=(f"=IFERROR("
                   f"FILTER("
                   f"CHOOSE({{1,2,3,4}},"
                   f"B${data_start}:B${data_start+11},"
                   f"H${data_start}:H${data_start+11},"
                   f"I${data_start}:I${data_start+11},"
                   f"J${data_start}:J${data_start+11}),"
                   f"J${data_start}:J${data_start+11}<>\"\"),"
                   f"\"아직 입력된 특이사항이 없습니다.\")"))


def _add_unclassified_section(ws, anchor_row):
    """뱅샐 대분류가 예산 카테고리에 없는 거래 자동 감지"""
    _section_title(ws, anchor_row, "⑤ 미분류 거래 감지  (예산 카테고리 미매핑 지출 자동 표시)")

    # 알려진 카테고리 배열 상수
    known_cats = sorted(set(main for main, *_ in BUDGET_ITEMS))
    cats_array = "{" + ",".join(f'"{c}"' for c in known_cats) + "}"

    hint_row = anchor_row + 1
    hint = ws.cell(row=hint_row, column=2,
                   value=f"※ 뱅샐 대분류가 예산 항목({', '.join(known_cats)})에 없는 지출 거래를 표시합니다. "
                         "해당 거래가 없으면 '미분류 거래 없음'이 표시됩니다.")
    hint.font = Font(size=8, color="888888", italic=True)
    hint.alignment = _align(h="left")
    ws.merge_cells(f"B{hint_row}:J{hint_row}")
    ws.row_dimensions[hint_row].height = 14

    # 헤더
    hrow = anchor_row + 2
    ws.row_dimensions[hrow].height = 22
    headers = ["날짜", "대분류", "소분류", "내용", "금액"]
    hcolors = ["1A2B4A", "C0392B", "C0392B", "2471A3", "C0392B"]
    for ci, (h, hc) in enumerate(zip(headers, hcolors), start=2):
        c = ws.cell(row=hrow, column=ci, value=h)
        c.fill = _fill(hc)
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.alignment = _align()
        c.border = _border()

    # FILTER: 미분류 거래 (대분류가 known_cats에 없는 지출)
    data_row = anchor_row + 3
    ws.cell(row=data_row, column=2,
            value=(f'=IFERROR('
                   f'FILTER('
                   f'CHOOSE({{1,2,3,4,5}},'
                   f'뱅샐입력!A$4:A${DATA_LAST_ROW},'
                   f'뱅샐입력!D$4:D${DATA_LAST_ROW},'
                   f'뱅샐입력!E$4:E${DATA_LAST_ROW},'
                   f'뱅샐입력!F$4:F${DATA_LAST_ROW},'
                   f'뱅샐입력!G$4:G${DATA_LAST_ROW}),'
                   f'(뱅샐입력!C$4:C${DATA_LAST_ROW}="지출")*'
                   f'(YEAR(뱅샐입력!A$4:A${DATA_LAST_ROW})=설정!C3)*'
                   f'(뱅샐입력!D$4:D${DATA_LAST_ROW}<>"")*'
                   f'ISERROR(MATCH(뱅샐입력!D$4:D${DATA_LAST_ROW},{cats_array},0))),'
                   f'"미분류 거래 없음")'))
    ws.cell(row=data_row, column=2).number_format = "YYYY-MM-DD"


def _section_title(ws, row, title):
    c = ws.cell(row=row, column=2, value=title)
    c.font = _font(bold=True, size=12, color="FFFFFF")
    c.fill = _fill("1A2B4A")
    c.alignment = _align(h="left")
    c.border = _border()
    ws.merge_cells(f"B{row}:J{row}")
    ws.row_dimensions[row].height = 24
