"""자산현황 시트 생성 모듈"""

from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter


def _side():
    return Side(style="thin", color="E2E8F0")

def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, size=11, color="000000"):
    return Font(name="Calibri", bold=bold, size=size, color=color)

def _align(h="center", v="center"):
    return Alignment(horizontal=h, vertical=v)


# (항목명, 기본값, 비고, 카테고리)  카테고리: asset/liability
ASSET_ITEMS = [
    # ── 자산 ──────────────────────────────────────────────────
    ("부동산",      900_000_000, "기흥역센트럴푸르지오",     "asset"),
    ("계좌(KB)",              0, "직접 입력",                "asset"),
    ("계좌(토스)",            0, "직접 입력",                "asset"),
    ("계좌(기타)",            0, "직접 입력",                "asset"),
    ("주식",                  0, "삼성증권 등",              "asset"),
    ("코인",                  0, "직접 입력",                "asset"),
    ("퇴직금(예상)",          0, "직접 입력",                "asset"),
    ("현금",                  0, "직접 입력",                "asset"),
    # ── 부채 ──────────────────────────────────────────────────
    ("대출잔액",              0, "디딤돌대출 등",            "liability"),
    ("기타부채",              0, "직접 입력",                "liability"),
]

# 월별 추이 테이블 컬럼 (B=월, C=부동산, D=계좌합계, E=주식, F=코인, G=퇴직금, H=대출잔액, I=순자산)
TREND_HEADERS  = ["월", "부동산", "계좌합계", "주식", "코인", "퇴직금", "대출잔액", "순자산"]
TREND_COLORS   = ["1A2B4A", "2471A3", "1E8449", "2471A3", "2471A3", "2471A3", "C0392B", "1A2B4A"]


def build(wb, year: int = 2026):
    ws = wb.create_sheet("자산현황")
    ws.sheet_view.showGridLines = False

    # A:여백  B:항목/월  C:금액/부동산  D:비고/계좌  E:주식  F:코인  G:퇴직금  H:대출  I:순자산  J:여백
    col_widths = {1: 4, 2: 14, 3: 18, 4: 14, 5: 14, 6: 14, 7: 14, 8: 14, 9: 14, 10: 4}
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for r in range(1, 80):
        ws.row_dimensions[r].height = 22

    # ── 제목 ─────────────────────────────────────────────────
    title = ws.cell(row=1, column=2, value="자산 현황")
    title.font = Font(name="Calibri Light", bold=False, size=14, color="1A2B4A")
    title.alignment = _align(h="left")
    ws.merge_cells("B1:I1")
    ws.row_dimensions[1].height = 30

    note = ws.cell(row=2, column=2,
                   value="※ 계좌 잔액, 주식, 코인, 퇴직금은 직접 입력하세요.")
    note.font = Font(size=9, color="888888", italic=True)
    note.alignment = _align(h="left")
    ws.merge_cells("B2:I2")

    # ── 스냅샷 헤더 ──────────────────────────────────────────
    header_row = 4
    for ci, (h, hc) in enumerate(zip(["항목", "금액", "비고"],
                                     ["1A2B4A", "1E8449", "1A2B4A"]), start=2):
        c = ws.cell(row=header_row, column=ci, value=h)
        c.fill = _fill(hc)
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.alignment = _align()
        c.border = _border()

    # ── 자산 항목 ─────────────────────────────────────────────
    row = 5
    asset_rows, liability_rows = [], []

    _sub_header(ws, row, "[ 자산 ]", "1E8449")
    row += 1
    for item, default, note_text, cat in ASSET_ITEMS:
        if cat != "asset":
            continue
        c_item = ws.cell(row=row, column=2, value=item)
        c_amt  = ws.cell(row=row, column=3, value=default)
        c_note = ws.cell(row=row, column=4, value=note_text)

        fill_c = "E8F8F0" if len(asset_rows) % 2 == 0 else "FFFFFF"
        for c in (c_item, c_amt, c_note):
            c.fill = _fill(fill_c)
            c.alignment = _align()
            c.border = _border()
        c_amt.number_format = '#,##0"원"'
        asset_rows.append(row)
        row += 1

    total_asset_row = row
    _total_row(ws, row, asset_rows, "총 자산", "1E8449")
    row += 2

    _sub_header(ws, row, "[ 부채 ]", "C0392B")
    row += 1
    for item, default, note_text, cat in ASSET_ITEMS:
        if cat != "liability":
            continue
        c_item = ws.cell(row=row, column=2, value=item)
        c_amt  = ws.cell(row=row, column=3, value=default)
        c_note = ws.cell(row=row, column=4, value=note_text)

        fill_c = "FDF2F2" if len(liability_rows) % 2 == 0 else "FFFFFF"
        for c in (c_item, c_amt, c_note):
            c.fill = _fill(fill_c)
            c.alignment = _align()
            c.border = _border()
        c_amt.number_format = '#,##0"원"'
        liability_rows.append(row)
        row += 1

    total_liability_row = row
    _total_row(ws, row, liability_rows, "총 부채", "C0392B")
    row += 2

    # ── 순자산 ───────────────────────────────────────────────
    for ci in range(2, 5):
        c = ws.cell(row=row, column=ci)
        c.fill = _fill("1A2B4A")
        c.font = _font(bold=True, color="FFFFFF", size=12)
        c.alignment = _align()
        c.border = _border()
    ws.cell(row=row, column=2, value="순 자산 (자산 - 부채)")
    net = ws.cell(row=row, column=3,
                  value=f"=C{total_asset_row}-C{total_liability_row}")
    net.number_format = '#,##0"원"'
    net.font = _font(bold=True, size=12, color="FFFFFF")
    row += 2

    # ════════════════════════════════════════════════════════
    # 월별 자산 추이 (매월 직접 입력 → 순자산 자동 계산)
    # ════════════════════════════════════════════════════════
    _section_title(ws, row, "월별 자산 추이  (매월 말 업데이트)")
    row += 1

    hint = ws.cell(row=row, column=2,
                   value="※ 부동산·계좌합계·주식·코인·퇴직금은 매월 직접 입력하세요. 대출잔액은 대출상환 시트에서 자동 참조됩니다.")
    hint.font = Font(size=9, color="888888", italic=True)
    hint.alignment = _align(h="left")
    ws.merge_cells(f"B{row}:I{row}")
    ws.row_dimensions[row].height = 14
    row += 1

    # 추이 테이블 헤더
    trend_header_row = row
    for ci, (h, hc) in enumerate(zip(TREND_HEADERS, TREND_COLORS), start=2):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = _fill(hc)
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.alignment = _align()
        c.border = _border()
    ws.row_dimensions[row].height = 22
    row += 1

    # 대출상환 스케줄 연계 상수
    # 대출 실행: 2026-03-30 / 첫 납입: 2026-04-30 (회차1 = 대출상환!H67)
    _LOAN_ISSUE_YEAR  = 2026
    _LOAN_ISSUE_MONTH = 3     # 3월: 대출 실행, 잔액 = 원금
    _LOAN_PAY_MONTH   = 4     # 4월: 첫 납입 (회차 1)
    _LOAN_SCHED_START = 67    # 대출상환 시트 스케줄 데이터 시작 행

    # 1~12월 데이터 행
    trend_data_start = row
    for month in range(1, 13):
        fill_c = "F4F8FC" if month % 2 == 0 else "FFFFFF"
        ws.row_dimensions[row].height = 21

        # B: 월 라벨
        c = ws.cell(row=row, column=2, value=f"{month}월")
        c.fill = _fill(fill_c)
        c.font = _font(bold=True)
        c.alignment = _align()
        c.border = _border()

        # 대출잔액 수식 계산 (H열 = col 8)
        # months_offset: 0 = 첫 납입(2026-04), 음수 = 이전, 양수 = 이후
        months_offset = (year - _LOAN_ISSUE_YEAR) * 12 + (month - _LOAN_PAY_MONTH)
        if months_offset < -1:
            # 대출 실행 전: 0
            loan_formula = 0
        elif months_offset == -1:
            # 대출 실행 월(2026-03): 원금 전액
            loan_formula = "='대출상환'!$C$5"
        else:
            # 첫 납입 이후: 해당 회차 잔액 참조
            sched_row = _LOAN_SCHED_START + months_offset
            loan_formula = f"=IFERROR('대출상환'!$H${sched_row},0)"

        # C~G: 수동 입력 (부동산, 계좌합계, 주식, 코인, 퇴직금)
        for ci in range(3, 8):
            c = ws.cell(row=row, column=ci, value=0)
            c.fill = _fill(fill_c)
            c.alignment = _align()
            c.border = _border()
            c.number_format = '#,##0"원"'

        # H: 대출잔액 (대출상환 시트 자동 참조)
        hc = ws.cell(row=row, column=8, value=loan_formula)
        hc.fill = _fill("F4F8FC")   # 자동참조 셀은 연한 파랑으로 구분
        hc.font = _font(color="1A2B4A")
        hc.alignment = _align()
        hc.border = _border()
        hc.number_format = '#,##0"원"'

        # I: 순자산 = (부동산+계좌+주식+코인+퇴직금) - 대출잔액
        net_c = ws.cell(row=row, column=9,
                        value=f"=C{row}+D{row}+E{row}+F{row}+G{row}-H{row}")
        net_c.fill = _fill("D6EAF8" if month % 2 == 0 else "F4F8FC")
        net_c.font = _font(bold=True, color="1A2B4A")
        net_c.alignment = _align()
        net_c.border = _border()
        net_c.number_format = '#,##0"원"'

        row += 1

    trend_data_end = row - 1

    # ── 추이 차트: 순자산 꺾은선 ─────────────────────────────
    chart_anchor_row = row + 1
    _add_trend_chart(ws, trend_data_start, trend_data_end, chart_anchor_row)

    return ws


def _add_trend_chart(ws, data_start, data_end, anchor_row):
    """월별 순자산 추이 꺾은선 차트"""
    chart = LineChart()
    chart.title = "월별 순자산 추이"
    chart.y_axis.title = "순자산 (원)"
    chart.x_axis.title = "월"
    chart.style = 2
    chart.width = 22
    chart.height = 12
    chart.y_axis.majorGridlines = None

    # 순자산 (I열 = col 9)
    net_ref = Reference(ws, min_col=9, min_row=data_start, max_row=data_end)
    chart.add_data(net_ref, titles_from_data=False)
    chart.series[0].title = SeriesLabel(v="순자산")
    chart.series[0].graphicalProperties.line.solidFill = "2471A3"
    chart.series[0].graphicalProperties.line.width = 28000  # 2.2pt
    chart.series[0].smooth = True

    # X축: 월 라벨 (B열 = col 2)
    cats = Reference(ws, min_col=2, min_row=data_start, max_row=data_end)
    chart.set_categories(cats)

    ws.add_chart(chart, f"B{anchor_row}")


def _section_title(ws, row, title):
    c = ws.cell(row=row, column=2, value=title)
    c.font = _font(bold=True, size=12, color="FFFFFF")
    c.fill = _fill("1A2B4A")
    c.alignment = _align(h="left")
    c.border = _border()
    ws.merge_cells(f"B{row}:I{row}")
    ws.row_dimensions[row].height = 24


def _sub_header(ws, row, title, color):
    c = ws.cell(row=row, column=2, value=title)
    c.font = _font(bold=True, color="FFFFFF")
    c.fill = _fill(color)
    c.alignment = _align(h="left")
    c.border = _border()
    ws.merge_cells(f"B{row}:D{row}")


def _total_row(ws, row, data_rows, label, color):
    c_label = ws.cell(row=row, column=2, value=label)
    c_label.font = _font(bold=True, color="FFFFFF")
    c_label.fill = _fill(color)
    c_label.alignment = _align()
    c_label.border = _border()

    row_refs = "+".join(f"C{r}" for r in data_rows)
    c_total = ws.cell(row=row, column=3, value=f"={row_refs}")
    c_total.number_format = '#,##0"원"'
    c_total.font = _font(bold=True)
    c_total.fill = _fill("D6EAF8")
    c_total.alignment = _align()
    c_total.border = _border()

    c_note = ws.cell(row=row, column=4, value="")
    c_note.fill = _fill("D6EAF8")
    c_note.border = _border()
