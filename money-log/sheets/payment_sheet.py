"""결제수단별 지출 집계 시트 생성 모듈"""

from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from .banksalad_sheet import DATA_LAST_ROW

MAX_PAYMENT_METHODS = 30   # UNIQUE 스필 대비 사전 할당 행 수


def _side():
    return Side(style="thin", color="E2E8F0")
def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)
def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)
def _font(bold=False, size=11, color="000000"):
    return Font(name="Calibri", bold=bold, size=size, color=color)
def _align(h="center", v="center"):
    return Alignment(horizontal=h, vertical=v, wrap_text=False)


def build(wb):
    ws = wb.create_sheet("결제수단")
    ws.sheet_view.showGridLines = False

    # A: 결제수단, B: 연간합계, C~N: 1~12월
    col_widths = {1: 26, 2: 14}
    for m in range(1, 13):
        col_widths[m + 2] = 12
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    # ── 제목 ─────────────────────────────────────────────────
    ws.row_dimensions[1].height = 28
    t = ws.cell(row=1, column=1, value='=설정!C3&"년 결제수단별 지출"')
    t.font = Font(name="Calibri Light", bold=False, size=14, color="1A2B4A")
    t.alignment = _align(h="left")
    ws.merge_cells("A1:N1")

    note = ws.cell(row=2, column=1,
                   value="※ 뱅샐입력 붙여넣기 후 자동 집계 (지출 타입만, 이체 제외). "
                         "카드·계좌 추가 시 자동 반영.")
    note.font = Font(size=9, color="888888", italic=True)
    note.alignment = _align(h="left")
    ws.merge_cells("A2:N2")
    ws.row_dimensions[2].height = 14

    # ── 헤더 행 ──────────────────────────────────────────────
    header_row = 4
    ws.row_dimensions[header_row].height = 22
    headers = ["결제수단", "연간 합계"] + [f"{m}월" for m in range(1, 13)]
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=ci, value=h)
        c.fill = _fill("1A2B4A")
        c.font = _font(bold=True, color="FFFFFF", size=10)
        c.alignment = _align()
        c.border = _border()

    # ── A5: UNIQUE — 뱅샐의 결제수단 목록 자동 추출 ──────────
    data_start = 5
    ws.cell(row=data_start, column=1,
            value=(f'=IFERROR(UNIQUE(FILTER('
                   f'뱅샐입력!I$4:I${DATA_LAST_ROW},'
                   f'(뱅샐입력!I$4:I${DATA_LAST_ROW}<>"")*'
                   f'(뱅샐입력!C$4:C${DATA_LAST_ROW}="지출")*'
                   f'(YEAR(뱅샐입력!A$4:A${DATA_LAST_ROW})=설정!C3))),'
                   f'"")')).alignment = _align(h="left")

    # ── B5~N{MAX}: SUMPRODUCT 사전 할당 ──────────────────────
    for i in range(MAX_PAYMENT_METHODS):
        r = data_start + i
        fill_c = "F4F8FC" if i % 2 == 0 else "FFFFFF"
        ws.row_dimensions[r].height = 20

        # A열 스타일 (UNIQUE 스필 대상 행)
        c = ws.cell(row=r, column=1)
        c.fill = _fill(fill_c)
        c.alignment = _align(h="left")
        c.border = _border()

        # B: 연간 합계
        ws.cell(row=r, column=2,
                value=(f'=IFERROR(IF(A{r}="","",SUMPRODUCT('
                       f'(뱅샐입력!I$4:I${DATA_LAST_ROW}=A{r})*'
                       f'(YEAR(뱅샐입력!A$4:A${DATA_LAST_ROW})=설정!C3)*'
                       f'(뱅샐입력!C$4:C${DATA_LAST_ROW}="지출")*'
                       f'ABS(뱅샐입력!G$4:G${DATA_LAST_ROW}))),0)'
                       )).number_format = '#,##0"원"'

        # C~N: 월별 합계
        for m in range(1, 13):
            col = m + 2
            ws.cell(row=r, column=col,
                    value=(f'=IFERROR(IF(A{r}="","",SUMPRODUCT('
                           f'(뱅샐입력!I$4:I${DATA_LAST_ROW}=A{r})*'
                           f'(MONTH(뱅샐입력!A$4:A${DATA_LAST_ROW})={m})*'
                           f'(YEAR(뱅샐입력!A$4:A${DATA_LAST_ROW})=설정!C3)*'
                           f'(뱅샐입력!C$4:C${DATA_LAST_ROW}="지출")*'
                           f'ABS(뱅샐입력!G$4:G${DATA_LAST_ROW}))),0)'
                           )).number_format = '#,##0"원"'

        for ci in range(2, 15):
            ws.cell(row=r, column=ci).fill = _fill(fill_c)
            ws.cell(row=r, column=ci).alignment = _align()
            ws.cell(row=r, column=ci).border = _border()

    # ── 합계 행 ──────────────────────────────────────────────
    total_r = data_start + MAX_PAYMENT_METHODS
    ws.row_dimensions[total_r].height = 22
    for ci in range(1, 15):
        c = ws.cell(row=total_r, column=ci)
        c.fill = _fill("1A2B4A")
        c.font = _font(bold=True, color="FFFFFF")
        c.alignment = _align()
        c.border = _border()
    ws.cell(row=total_r, column=1, value="합계")
    for ci in range(2, 15):
        col_l = get_column_letter(ci)
        ws.cell(row=total_r, column=ci,
                value=f"=SUM({col_l}{data_start}:{col_l}{total_r-1})").number_format = '#,##0"원"'

    ws.freeze_panes = "B5"
    return ws
